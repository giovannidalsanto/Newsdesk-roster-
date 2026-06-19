#!/usr/bin/env python3
"""
Convert a filled EBS Newsdesk roster .xlsx into roster.json for the web app.
Usage: python convert.py roster.xlsx   (defaults to roster.xlsx -> roster.json)
"""
import sys, json, datetime
from openpyxl import load_workbook

ROLES = ["MMC Frontpage","EC Hotline","Octopus – Daily","Octopus – Week Ahead","Events",
  "Mission Coordination","Watch Party","TXM with Francis","Pre-production","Production",
  "Troubleshooting","Planning","Agencies","INS Follow-up"]
PEOPLE = ["Giovanni","Zoe","Claudio"]
WEEKDAY_BANDS = ["Early","Morning","Lunch","Afternoon","Late"]
DAYORDER = ["Mon","Tue","Wed","Thu","Fri","Sat","Sun"]

def parse_week_start(v):
    if isinstance(v,(datetime.datetime,datetime.date)):
        return datetime.date(v.year,v.month,v.day)
    s=str(v).strip()
    for fmt in ("%Y-%m-%d","%d/%m/%Y","%d-%m-%Y","%m/%d/%Y"):
        try: return datetime.datetime.strptime(s,fmt).date()
        except: pass
    raise SystemExit(f"Could not read WEEK STARTING date: {v!r}. Use YYYY-MM-DD.")

def main():
    path = sys.argv[1] if len(sys.argv)>1 else "roster.xlsx"
    out = sys.argv[2] if len(sys.argv)>2 else "roster.json"
    wb = load_workbook(path, data_only=True)
    rd = wb["READ ME FIRST"]
    monday = parse_week_start(rd["C3"].value)
    days=[]
    for di,tab in enumerate(DAYORDER):
        if tab not in wb.sheetnames: continue
        ws=wb[tab]; weekend = tab in ("Sat","Sun")
        bands = ["All day"] if weekend else WEEKDAY_BANDS
        status={}
        for i,p in enumerate(PEOPLE):
            v=ws.cell(row=4,column=2+i).value
            status[p]= (str(v).strip() if v else "Working")
        grid={}
        for i,role in enumerate(ROLES):
            r=7+i; cells={}
            for j,b in enumerate(bands):
                v=ws.cell(row=r,column=2+j).value
                cells[b]= (str(v).strip() if isinstance(v,str) and v.strip() else None)
            grid[role]=cells
        date = monday + datetime.timedelta(days=di)
        days.append({"day":tab,"date":date.isoformat(),"weekend":weekend,
                     "bands":bands,"status":status,"grid":grid})
    feed={"weekStart":monday.isoformat(),"roles":ROLES,"people":PEOPLE,
          "weekdayBands":WEEKDAY_BANDS,"generatedAt":datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
          "days":days}
    json.dump(feed, open(out,"w",encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"{out} written: week of {monday}, {sum(1 for d in days for r in ROLES for b in d['bands'] if d['grid'][r][b] in PEOPLE)} assignments")

if __name__=="__main__": main()
